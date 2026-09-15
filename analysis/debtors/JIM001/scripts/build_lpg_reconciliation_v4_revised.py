#!/usr/bin/env python3
"""Build a revised JIM001_LPG_Reconciliation_v4 workbook, scoped to
month_year >= 2023-03, sourced entirely from monthly_lpg_insights.csv
(not invoices.csv/payments.csv raw ledgers, not the Google Sheet).

Two sheets:
  - "Monthly LPG Insights (2023-03+)": one row per month, mirroring
    monthly_lpg_insights.csv's own columns, with a totals row (live
    SUM formulas).
  - "Summary": headline reconciliation figures + source/evidence note.
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[4]
CSV_PATH = ROOT / "analysis/debtors/JIM001/data/monthly_lpg_insights.csv"
OUT_PATH = ROOT / "analysis/debtors/JIM001/reports/JIM001_LPG_Reconciliation_v4_Revised_2023-03.xlsx"

PERIOD_START = "2023-03"

CURRENCY = 'R#,##0.00;[Red]-R#,##0.00;"—"'
FONT_NAME = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
TOTAL_FONT = Font(name=FONT_NAME, bold=True)
DATA_FONT = Font(name=FONT_NAME)
STATUS_FILL = {
    "FULLY_SETTLED": PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"),
    "PARTIALLY_SETTLED": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    "UNPAID": PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid"),
    "OVERPAID": PatternFill(start_color="CFE2F3", end_color="CFE2F3", fill_type="solid"),
}


def load_rows():
    with open(CSV_PATH, newline="") as f:
        rows = list(csv.DictReader(f))
    return [r for r in rows if r["month_year"] >= PERIOD_START]


def build_main_sheet(wb, rows):
    ws = wb.active
    ws.title = "Monthly LPG Insights (2023-03+)"

    headers = [
        "Year", "Month", "Month-Year", "Net LPG Invoiced (R)",
        "Payment Doc(s)", "Payment Date(s)", "Payment Allocated (R)",
        "Difference (R)", "Status", "Review Required", "Notes",
    ]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

    r = 2
    for row in rows:
        ws.cell(row=r, column=1, value=int(row["year"]))
        ws.cell(row=r, column=2, value=row["month"])
        ws.cell(row=r, column=3, value=row["month_year"])
        ws.cell(row=r, column=4, value=float(row["net_lpg_invoiced"]))
        ws.cell(row=r, column=5, value=row["payment_refs_allocated"])
        ws.cell(row=r, column=6, value=row["payment_dates"])
        ws.cell(row=r, column=7, value=float(row["payment_total_allocated"]))
        ws.cell(row=r, column=8, value=float(row["difference"]))
        ws.cell(row=r, column=9, value=row["month_status"])
        ws.cell(row=r, column=10, value=row["review_required"])
        ws.cell(row=r, column=11, value=row["notes"])

        fill = STATUS_FILL.get(row["month_status"])
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = DATA_FONT
            if fill:
                cell.fill = fill
            if c in (4, 7, 8):
                cell.number_format = CURRENCY
            if c == 11:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    last_data_row = r - 1
    total_row = r
    ws.cell(row=total_row, column=3, value="TOTAL")
    ws.cell(row=total_row, column=4, value=f"=SUM(D2:D{last_data_row})")
    ws.cell(row=total_row, column=7, value=f"=SUM(G2:G{last_data_row})")
    ws.cell(row=total_row, column=8, value=f"=SUM(H2:H{last_data_row})")
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=total_row, column=c)
        cell.font = TOTAL_FONT
        if c in (4, 7, 8):
            cell.number_format = CURRENCY

    widths = [6, 11, 11, 16, 14, 24, 16, 14, 17, 14, 70]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    return last_data_row, total_row


def build_summary_sheet(wb, rows, last_data_row, total_row, main_sheet_name):
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 20

    def title(r, text):
        cell = ws.cell(row=r, column=1, value=text)
        cell.font = Font(name=FONT_NAME, bold=True, size=13)

    def label_value(r, label, value, is_formula=False, currency=True):
        lc = ws.cell(row=r, column=1, value=label)
        lc.font = DATA_FONT
        vc = ws.cell(row=r, column=2, value=value)
        vc.font = TOTAL_FONT if is_formula else DATA_FONT
        if currency:
            vc.number_format = CURRENCY

    title(1, "JIM001 — LPG Reconciliation v4 (Revised, 2023-03 onward)")
    ws.cell(row=2, column=1, value=(
        "Source: monthly_lpg_insights.csv only (not invoices.csv/payments.csv raw ledgers, "
        "not JIM001_LPG_Reconciliation_v4's original Google-Sheet-derived structure)."
    )).font = DATA_FONT
    ws.cell(row=3, column=1, value=f"Period: {rows[0]['month_year']} through {rows[-1]['month_year']} ({len(rows)} months).").font = DATA_FONT
    ws.cell(row=4, column=1, value="Generated 2026-09-15.").font = DATA_FONT

    r = 6
    title(r, "Reconciliation totals")
    r += 1
    label_value(r, "Total net LPG invoiced", f"='{main_sheet_name}'!D{total_row}", is_formula=True)
    r += 1
    label_value(r, "Total payment allocated", f"='{main_sheet_name}'!G{total_row}", is_formula=True)
    r += 1
    label_value(r, "Total difference (net unpaid)", f"='{main_sheet_name}'!H{total_row}", is_formula=True)
    r += 2

    title(r, "Months by status")
    r += 1
    from collections import Counter
    counts = Counter(row["month_status"] for row in rows)
    for status in ("FULLY_SETTLED", "PARTIALLY_SETTLED", "UNPAID", "OVERPAID"):
        if counts.get(status):
            label_value(r, status, counts[status], is_formula=False, currency=False)
            r += 1
    r += 1

    title(r, "Evidence note")
    r += 1
    note = (
        "Figures here are exactly as recorded in monthly_lpg_insights.csv as of 2026-09-15 -- "
        "including the resolved 2025-01/02 (docs 38846/39812), 2025-07 (doc 40746, PROVEN Tier 2), "
        "and 2025-12 (doc 44555, account-owner decision, not independently proven) entries. "
        "See JIM001_Exact_Sum_Bridge_Review_2026-09-13.md for full evidence-tier detail per month; "
        "do not treat this workbook as a higher evidence tier than that report."
    )
    cell = ws.cell(row=r, column=1, value=note)
    cell.font = DATA_FONT
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws.row_dimensions[r].height = 60


def main():
    rows = load_rows()
    wb = Workbook()
    last_data_row, total_row = build_main_sheet(wb, rows)
    build_summary_sheet(wb, rows, last_data_row, total_row, wb.sheetnames[0])
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} -- {len(rows)} months ({rows[0]['month_year']} to {rows[-1]['month_year']})")


if __name__ == "__main__":
    main()
