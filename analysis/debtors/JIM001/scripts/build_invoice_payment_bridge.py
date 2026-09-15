#!/usr/bin/env python3
"""JIM001 invoice-vs-payment bridge: payment original amount vs allocated amount.

Allocation evidence, strongest first:
  ERP screen  - ERP View Payment Allocations screen supplied by the account owner
  interlock   - cent-exact arithmetic interlock with an ERP screen
  proven      - invoice-level match established and ratified in the Bridge review
  alloc_edges - allocation_edges.csv's own allocated_amount
  none        - no allocation record anywhere
"""
import csv, json
from collections import defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[4]
D = ROOT / "analysis/debtors/JIM001/data"
OUT = ROOT / "analysis/debtors/JIM001/reports/JIM001_Invoice_vs_Payment_Bridge.xlsx"

CUR = 'R#,##0.00;[Red]-R#,##0.00;"—"'
F = "Arial"
HF = Font(name=F, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", start_color="1F4E78")
BOLD = Font(name=F, bold=True)
BODY = Font(name=F)
TITLE = Font(name=F, bold=True, size=13)
FILLS = {"ERP screen": PatternFill("solid", start_color="D9EAD3"),
         "interlock": PatternFill("solid", start_color="D9EAD3"),
         "proven": PatternFill("solid", start_color="D9EAD3"),
         "none": PatternFill("solid", start_color="F4CCCC")}

ERP = {"33810", "34425", "35270", "36139", "36988"}
ILOCK = {"32896"}
PROVEN = {"38846", "39812"}
ORPHAN = {"4832","4839","4807","4898","4841","5043","5047","5226",
          "6381","7216","7226","7348","7489","7572"}

pays = list(csv.DictReader(open(D/"payments.csv", newline='')))
edges = list(csv.DictReader(open(D/"allocation_edges.csv", newline='')))
inv = [r for r in csv.DictReader(open(D/"invoices.csv", newline='')) if r['debtor_code'] == 'JIM001']
dm = json.load(open(D/"dashboard_metrics.json"))

alloc, ntgt = defaultdict(float), defaultdict(int)
for e in edges:
    if e['allocation_type'] != 'UNALLOCATED_PORTION':
        d = e['payment_doc'].lstrip('0')
        alloc[d] += float(e['allocated_amount'] or 0)
        ntgt[d] += 1

rows = []
for p in sorted(pays, key=lambda x: x['payment_date']):
    d = p['payment_doc'].lstrip('0')
    g = abs(float(p['amount']))
    if d in ERP:      a, ev = g, "ERP screen"
    elif d in ILOCK:  a, ev = g, "interlock"
    elif d in PROVEN: a, ev = g, "proven"
    else:
        a = alloc.get(d, 0.0)
        ev = "alloc_edges" if a else "none"
    rows.append({"date": p['payment_date'], "doc": d, "batch": p['batch_ref'],
                 "gross": g, "alloc": a, "n": ntgt.get(d, 0), "ev": ev,
                 "orphan": d in ORPHAN})

lpg = sum(float(r['amount_incl']) for r in inv if r['is_lpg'] == 'True')
cyl = sum(float(r['amount_incl']) for r in inv if r['is_cyl'] == 'True')
oth = sum(float(r['amount_incl']) for r in inv if r['is_lpg'] != 'True' and r['is_cyl'] != 'True')

wb = Workbook()

# ---------- Payments sheet ----------
ws = wb.active
ws.title = "Payments"
hdr = ["Payment date", "Doc #", "Batch ref", "ORIGINAL amount (R)",
       "ALLOCATED to invoices (R)", "UNALLOCATED (R)", "# invoices", "Allocation evidence"]
ws.append(hdr)
for c in range(1, len(hdr)+1):
    cell = ws.cell(row=1, column=c)
    cell.font, cell.fill = HF, HFILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.freeze_panes = "A2"

r = 2
for x in rows:
    ws.cell(row=r, column=1, value=x["date"])
    ws.cell(row=r, column=2, value=x["doc"])
    ws.cell(row=r, column=3, value=x["batch"])
    ws.cell(row=r, column=4, value=round(x["gross"], 2))
    ws.cell(row=r, column=5, value=round(x["alloc"], 2))
    ws.cell(row=r, column=6, value=f"=D{r}-E{r}")
    ws.cell(row=r, column=7, value=x["n"])
    ws.cell(row=r, column=8, value=x["ev"] + (" (pre-2021 orphan)" if x["orphan"] else ""))
    for c in range(1, len(hdr)+1):
        cell = ws.cell(row=r, column=c)
        cell.font = BODY
        if c in (4, 5, 6):
            cell.number_format = CUR
        if x["ev"] in FILLS:
            cell.fill = FILLS[x["ev"]]
    r += 1

last, tr = r-1, r
ws.cell(row=tr, column=3, value="TOTAL")
for col in ("D", "E", "F"):
    ws.cell(row=tr, column={"D":4,"E":5,"F":6}[col], value=f"=SUM({col}2:{col}{last})")
for c in range(1, len(hdr)+1):
    cell = ws.cell(row=tr, column=c)
    cell.font = BOLD
    if c in (4, 5, 6):
        cell.number_format = CUR
for i, w in enumerate([13, 9, 11, 19, 22, 16, 11, 30], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

# ---------- Bridge sheet ----------
bs = wb.create_sheet("Bridge", 0)
bs.column_dimensions["A"].width = 52
bs.column_dimensions["B"].width = 18

def title(rw, txt):
    c = bs.cell(row=rw, column=1, value=txt); c.font = TITLE

def line(rw, label, val, bold=False, cur=True):
    lc = bs.cell(row=rw, column=1, value=label); lc.font = BOLD if bold else BODY
    vc = bs.cell(row=rw, column=2, value=val);   vc.font = BOLD if bold else BODY
    if cur: vc.number_format = CUR

title(1, "JIM001 — Invoice vs Payment Bridge")
bs.cell(row=2, column=1, value="Period 2018-12 to 2026-06 · generated 2026-09-15 · source: invoices.csv, payments.csv, allocation_edges.csv, ERP allocation screens").font = BODY

title(4, "Invoice side — what was billed")
line(5, "LPG (gas)", round(lpg, 2))
line(6, "CYL (cylinder deposits, net of -EMPTY credit notes)", round(cyl, 2))
line(7, "OTHER (interest, net of same-day reversals)", round(oth, 2))
line(8, "Total billed", "=SUM(B5:B7)", bold=True)

title(10, "Payment side — what was received, and how much is allocated")
line(11, "Payments received (77 documents, gross)", f"=Payments!D{tr}", bold=True)
line(12, "   of which ALLOCATED to specific invoices", f"=Payments!E{tr}")
line(13, "   of which UNALLOCATED", f"=Payments!F{tr}")
line(14, "   allocated %", f"=B12/B11", cur=False)
bs.cell(row=14, column=2).number_format = "0.0%"

title(16, "Bridge to the ERP balance")
line(17, "Total billed", "=B8")
line(18, "less payments received", "=-B11")
line(19, "= Closing balance, computed", "=B17+B18", bold=True)
line(20, "add pre-2018 opening journal entry (documented, static)", 7350.00)
line(21, "add ERP residual variance (dashboard_metrics.erp_residual_variance)", dm["erp_residual_variance"])
line(22, "= Reconciled balance", "=B19+B20+B21", bold=True)
line(23, "ERP stated balance (dashboard_metrics)", dm["corrected_erp_stated_balance"], bold=True)
line(24, "VARIANCE", "=B22-B23", bold=True)

title(26, "The unallocated cash, broken down")
orph = sum(round(x["gross"]-x["alloc"], 2) for x in rows if x["orphan"] and x["gross"]-x["alloc"] > 0.01)
other = sum(round(x["gross"]-x["alloc"], 2) for x in rows if not x["orphan"] and x["gross"]-x["alloc"] > 0.01)
line(27, "Pre-2021 orphan payments (14 docs, unresolvable in-system)", round(orph, 2))
line(28, "All other gaps", round(other, 2))
line(29, "Total unallocated", "=SUM(B27:B28)", bold=True)
line(31, "ERP's own 'unmatched_overpayments' field, for comparison", dm["unmatched_overpayments"])
line(32, "difference vs our figure", "=B31-B29")

note = ("Reading this: ORIGINAL is the payment's gross as banked. ALLOCATED is how much of it the "
        "evidence ties to specific invoices. A gap is NOT necessarily missing money — for most docs "
        "it means the allocation record is incomplete, not that cash went astray. The five 2024 "
        "receipts marked 'ERP screen' are 100% allocated per the ERP's own View Payment Allocations "
        "screen; doc 32896 is proven by a cent-exact interlock with one of those screens; docs 38846 "
        "and 39812 are proven against Jan/Feb 2025 invoices. allocation_edges.csv is stale for the "
        "last two and still tags them UNALLOCATED — corrected here. The pre-2021 orphans are the one "
        "block that is genuinely unresolved and needs external bank records.")
c = bs.cell(row=34, column=1, value=note)
c.font = BODY; c.alignment = Alignment(wrap_text=True, vertical="top")
bs.merge_cells(start_row=34, start_column=1, end_row=34, end_column=2)
bs.row_dimensions[34].height = 105

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"Wrote {OUT}")
print(f"  billed R{lpg+cyl+oth:,.2f} | received R{sum(x['gross'] for x in rows):,.2f} | "
      f"allocated R{sum(x['alloc'] for x in rows):,.2f} | unallocated R{sum(x['gross']-x['alloc'] for x in rows):,.2f}")
